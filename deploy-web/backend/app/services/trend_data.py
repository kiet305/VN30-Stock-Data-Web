from __future__ import annotations

import csv
import os
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.schemas import TrendItem, TrendResponse


TREND_TIMEZONE = timezone(timedelta(hours=7))
TREND_FILE_PATTERN = "trending_*.csv"
TREND_ALIASES = {
    "vssid": ["bảo hiểm xã hội việt nam"],
    "xăng e10": ["xăng sinh học"],
    "bạc": ["giá bạc phú quý"],
    "cà phê": ["giá cà"],
}


def _trend_data_dir() -> Path:
    configured = os.getenv("TREND_DATA_DIR")
    if configured:
        return Path(configured)

    container_path = Path("/app/trend")
    if container_path.exists():
        return container_path

    return Path(__file__).resolve().parents[4] / "trend"


def _selected_trend_file() -> Path:
    configured_file = os.getenv("TREND_DATA_FILE")
    if configured_file:
        trend_file = Path(configured_file)
        if not trend_file.is_absolute():
            trend_file = _trend_data_dir() / trend_file
        if not trend_file.exists():
            raise FileNotFoundError(f"Không tìm thấy file xu hướng: {trend_file}")
        return trend_file

    files = sorted(
        _trend_data_dir().glob(TREND_FILE_PATTERN),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    if not files:
        raise FileNotFoundError(f"Không tìm thấy file {TREND_FILE_PATTERN} trong thư mục trend")

    return files[0]


def _selected_sparkline_file() -> Path | None:
    configured_file = os.getenv("TREND_SPARKLINE_FILE")
    if configured_file:
        sparkline_file = Path(configured_file)
        if not sparkline_file.is_absolute():
            sparkline_file = _trend_data_dir() / sparkline_file
        if not sparkline_file.exists():
            raise FileNotFoundError(f"Không tìm thấy file đường xu hướng: {sparkline_file}")
        return sparkline_file

    sparkline_file = _trend_data_dir() / "trends (2).xlsx"
    return sparkline_file if sparkline_file.exists() else None


def _normalize_keyword(value: str) -> str:
    return unicodedata.normalize("NFC", value).strip().casefold()


def _load_sparkline_points() -> dict[str, list[float]]:
    sparkline_file = _selected_sparkline_file()
    if sparkline_file is None:
        return {}

    workbook = load_workbook(sparkline_file, data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return {}

    keywords = [str(header).strip() if header is not None else "" for header in headers]
    points_by_keyword: dict[str, list[float]] = {
        _normalize_keyword(keyword): [] for keyword in keywords[1:] if keyword
    }

    for row in rows:
        for index, keyword in enumerate(keywords[1:], start=1):
            if not keyword:
                continue

            value = row[index] if index < len(row) else None
            if isinstance(value, (int, float)):
                points_by_keyword[_normalize_keyword(keyword)].append(float(value))

    workbook.close()
    return points_by_keyword


def _match_sparkline_points(row: dict[str, str], points_by_keyword: dict[str, list[float]]) -> list[float]:
    keyword = (row.get("Trends") or "").strip()
    candidates = [keyword, *_split_breakdown(row.get("Trend breakdown") or "")]
    candidates.extend(TREND_ALIASES.get(_normalize_keyword(keyword), []))

    for candidate in candidates:
        points = points_by_keyword.get(_normalize_keyword(candidate))
        if points:
            return points[-7:]

    return []


def _parse_trend_datetime(value: str | None) -> datetime | None:
    if not value:
        return None

    clean_value = value.replace("\u202f", " ").replace("\xa0", " ").strip()
    if " UTC" in clean_value:
        clean_value = clean_value.split(" UTC", 1)[0]

    try:
        parsed = datetime.strptime(clean_value, "%b %d, %Y at %I:%M:%S %p")
    except ValueError:
        return None

    return parsed.replace(tzinfo=TREND_TIMEZONE)


def _relative_label(value: datetime | None, reference: datetime | None) -> str:
    if value is None or reference is None:
        return "--"

    delta = reference - value
    if delta.total_seconds() < 0:
        return value.strftime("%d/%m/%Y")

    minutes = int(delta.total_seconds() // 60)
    if minutes < 60:
        return f"{max(minutes, 1)} phút trước"

    hours = minutes // 60
    if hours < 24:
        return f"{hours} giờ trước"

    if value.date() == (reference - timedelta(days=1)).date():
        return "Hôm qua"

    days = max(delta.days, 1)
    return f"{days} ngày trước"


def _duration_label(started: datetime | None, ended: datetime | None) -> str | None:
    if started is None or ended is None:
        return None

    total_minutes = max(int((ended - started).total_seconds() // 60), 0)
    hours = total_minutes // 60
    days = hours // 24
    remaining_hours = hours % 24

    if days and remaining_hours:
        return f"Kéo dài {days} ngày {remaining_hours} giờ"
    if days:
        return f"Kéo dài {days} ngày"
    if hours:
        return f"Kéo dài {hours} giờ"
    return f"Kéo dài {max(total_minutes, 1)} phút"


def _explore_url(raw_url: str) -> str:
    if raw_url.startswith("http://") or raw_url.startswith("https://"):
        return raw_url
    if raw_url.startswith("./"):
        return f"https://trends.google.com/trends/{raw_url[2:]}"
    return raw_url


def _split_breakdown(value: str) -> list[str]:
    return [item.strip() for item in value.split(",") if item.strip()]


def get_trends(limit: int = 50) -> TrendResponse:
    trend_file = _selected_trend_file()
    updated_at = datetime.fromtimestamp(trend_file.stat().st_mtime, tz=TREND_TIMEZONE)
    points_by_keyword = _load_sparkline_points()

    with trend_file.open(encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.DictReader(csv_file))

    items: list[TrendItem] = []
    for index, row in enumerate(rows[:limit], start=1):
        started = _parse_trend_datetime(row.get("Started"))
        ended = _parse_trend_datetime(row.get("Ended"))
        active = ended is None

        items.append(
            TrendItem(
                rank=index,
                keyword=(row.get("Trends") or "").strip(),
                search_volume=(row.get("Search volume") or "--").strip(),
                started=started,
                ended=ended,
                started_label=_relative_label(started, updated_at),
                status_label="Đang hoạt động" if active else "Đã kết thúc",
                duration_label=None if active else _duration_label(started, ended),
                active=active,
                trend_breakdown=_split_breakdown(row.get("Trend breakdown") or ""),
                trend_points=_match_sparkline_points(row, points_by_keyword),
                explore_url=_explore_url((row.get("Explore link") or "").strip()),
            )
        )

    return TrendResponse(
        source_file=trend_file.name,
        updated_at=updated_at,
        total=len(rows),
        items=items,
    )
